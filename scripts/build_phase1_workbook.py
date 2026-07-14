#!/usr/bin/env python3
"""Build the CVR Engineering Phase 1 baseline workbook.

Sources (see data/):
  - CVR_Rate_Sheet_2026.pdf            -> Rates tab
  - CVR_EPC_KPI_2026-07-11.pdf         -> Roster hours, Baseline P&L, EPC, checks

All published figures come from the KPI report (week ending 7/11/2026).
Cost of Sales and Operating Expenses on the Baseline P&L are DERIVED
(revenue - GP, GP - net income) because only Revenue / Gross Profit /
Net Income were published.

Output: model/CVR_Model_Phase1_Baseline.xlsx
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- styling ---
NAVY = "1F3864"
LIGHT = "D9E2F3"
INPUT_BLUE = "0070C0"
FLAG_ORANGE = "C55A11"
GREY = "F2F2F2"

F_TITLE = Font(bold=True, size=14, color=NAVY)
F_H = Font(bold=True, color="FFFFFF")
F_SUB = Font(bold=True, color=NAVY)
F_INPUT = Font(color=INPUT_BLUE)
F_FLAG = Font(color=FLAG_ORANGE, italic=True)
F_NOTE = Font(italic=True, size=9, color="595959")
FILL_H = PatternFill("solid", fgColor=NAVY)
FILL_SUB = PatternFill("solid", fgColor=LIGHT)
FILL_TOT = PatternFill("solid", fgColor=GREY)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

USD = '"$"#,##0.00'
K1 = '#,##0.0,"K"'          # values entered in dollars, shown as $K? no
KFMT = '#,##0.0'             # figures are entered in $K already
HRS = '#,##0.0'
PCT = '0.0%'


def header(ws, row, cols, labels):
    for c, label in zip(cols, labels):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = F_H
        cell.fill = FILL_H
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BOX


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_NOTE


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


wb = Workbook()

# ==================================================================== README
ws = wb.active
ws.title = "README"
title(ws, "CVR Engineering — Phase 1 Baseline Model",
      "Generated from scripts/build_phase1_workbook.py — do not hand-edit calc cells; blue cells are inputs.")
rows = [
    "",
    "PURPOSE",
    "Phase 1 of the financial model build: a quantified current-state baseline reconciled to the",
    "CVR & EPC KPI report (week ending 7/11/2026) and the 2026 rate sheet. It answers:",
    "  1. What did CVR actually earn and spend, January–May 2026?",
    "  2. What is the team's capacity worth at card rates, and what share of that is being realized?",
    "  3. What do utilization and realization have to become for the division's revenue targets?",
    "",
    "TABS",
    "  Assumptions        — global inputs (blue). Fringe load and overhead allocation are TBD from finance.",
    "  Rates              — 2026 rate sheet, effective January 2026.",
    "  Roster & Hours     — people, classifications, YTD hours and utilization vs target, value at card rates.",
    "  Baseline P&L       — CVR monthly actuals vs budget, Jan–May 2026. COGS/OpEx are derived lines.",
    "  HWC EPC Baseline   — EPC line financials for context (engineering outcomes affect this P&L).",
    "  Realization        — the $/hr gap analysis and the utilization x realization revenue grid.",
    "  Checks & Open Items— tie-outs to the KPI report and the questions finance must answer.",
    "",
    "COLOR KEY",
    "  Blue  = input / assumption you can change",
    "  Black = calculated or sourced from a document",
    "  Orange italics = assumption that needs confirmation",
    "",
    "SOURCES",
    "  data/CVR_Rate_Sheet_2026.pdf",
    "  data/CVR_EPC_KPI_2026-07-11.pdf",
]
for i, text in enumerate(rows, start=3):
    c = ws.cell(row=i, column=1, value=text)
    if text in ("PURPOSE", "TABS", "COLOR KEY", "SOURCES"):
        c.font = F_SUB
widths(ws, {"A": 110})

# ================================================================ Assumptions
ws = wb.create_sheet("Assumptions")
title(ws, "Global Assumptions", "Blue cells are inputs. TBD items block nothing but are flagged on Checks & Open Items.")
header(ws, 4, [1, 2, 3], ["Assumption", "Value", "Notes"])
assumptions = [
    ("Standard hours per FTE-year", 2080, "40 hrs x 52 wks; adjust for PTO/holiday policy if preferred"),
    ("Median target billable utilization", 0.75, "KPI report gauge target (median of per-person targets)"),
    ("Active billable headcount (current)", 10, "People with recorded hours in the KPI report"),
    ("Hours-to-May adjustment factor", 0.787, "KPI hours run through 7/11; factor scales them to Jan–May to match financials (21.6/27.4 wks)"),
    ("Fringe & benefit load on wages", "TBD", "REQUIRED from USC finance — drives loaded cost and true gross margin"),
    ("USC corporate overhead allocation", "TBD", "REQUIRED from USC finance — methodology and current $"),
    ("Interest expense basis (OILI)", "TBD", "REQUIRED from USC finance"),
    ("Year-2 divisional sales goal ($K)", 6000, "Divisional business target"),
    ("Year-2 OILI return target", 0.20, "OILI / revenue"),
]
r = 5
for name, val, note in assumptions:
    ws.cell(row=r, column=1, value=name).border = BOX
    v = ws.cell(row=r, column=2, value=val)
    v.border = BOX
    v.font = F_FLAG if val == "TBD" else F_INPUT
    if isinstance(val, float) and val <= 1:
        v.number_format = PCT
    elif isinstance(val, (int, float)):
        v.number_format = '#,##0'
    n = ws.cell(row=r, column=3, value=note)
    n.border = BOX
    n.font = F_NOTE
    r += 1
widths(ws, {"A": 40, "B": 14, "C": 80})

# ====================================================================== Rates
ws = wb.create_sheet("Rates")
title(ws, "CVR Engineering — Hourly Rates and Charges",
      "Effective January 2026. Source: data/CVR_Rate_Sheet_2026.pdf")
header(ws, 4, [1, 2], ["Classification", "Rate / hr"])
RATES = [
    ("Administrative Assistant", 59.62),
    ("Graphics Technician", 76.14),
    ("Project Manager", 159.86),
    ("Designer I", 79.91),
    ("Designer II", 95.70),
    ("Senior Designer", 105.01),
    ("Engineer I", 112.27),
    ("Engineer II", 138.09),
    ("Engineer III", 152.21),
    ("Engineer IV", 172.37),
    ("Senior Engineer", 241.57),
    ("Principal Engineer", 360.72),
]
r = 5
for name, rate in RATES:
    ws.cell(row=r, column=1, value=name).border = BOX
    c = ws.cell(row=r, column=2, value=rate)
    c.number_format = USD
    c.border = BOX
    r += 1
notes = [
    "Travel time in the interest of the work billed at full rates.",
    "No overtime multipliers — excess hours billed at standard rates.",
    "Mileage at IRS rate; ground/air travel and living expenses at cost.",
    "Purchases, out-of-scope subcontract work and other expenses at cost + 10%.",
]
r += 1
for note in notes:
    ws.cell(row=r, column=1, value="• " + note).font = F_NOTE
    r += 1
widths(ws, {"A": 30, "B": 12})
RATE_ROWS = {name: i + 5 for i, (name, _) in enumerate(RATES)}

# ============================================================ Roster & Hours
ws = wb.create_sheet("Roster & Hours")
title(ws, "Roster, YTD Hours and Utilization (2026 through week ending 7/11)",
      "Hours and splits from the KPI report utilization page. Orange classifications are assumptions pending confirmation.")
cols = list(range(1, 13))
header(ws, 4, cols, [
    "Name", "Billing classification", "Class source", "Rate / hr",
    "Total hrs YTD", "Billable %", "Non-billable %", "Bids %",
    "Billable hrs", "Target util %", "Util vs target", "Billable value @ card",
])
# name, class, source(confirmed/assumed), total hrs, billable%, nonbill%, bids%, target%
ROSTER = [
    ("Craig Peterson",   "Engineer I",          "org chart",              1017.00, 0.796, 0.204, 0.000, 0.754),
    ("Bhkti Patel",      "Engineer III",        "org chart",              1045.00, 0.744, 0.202, 0.055, 0.707),
    ("Taylor Nelson",    "Engineer II",         "ASSUMED — confirm",      1045.00, 0.610, 0.337, 0.054, 0.801),
    ("Hayley Worthen",   "Senior Engineer",     "org chart",               952.00, 0.669, 0.162, 0.169, 0.613),
    ("Francis Wagner",   "Engineer III",        "org chart",               992.00, 0.625, 0.343, 0.032, 0.707),
    ("Collin Allen",     "Graphics Technician", "org chart",               985.25, 0.576, 0.424, 0.000, 0.801),
    ("Kevin Metts",      "Designer II",         "ASSUMED — confirm",       890.00, 0.384, 0.598, 0.018, 0.754),
    ("Roy Pierce",       "Principal Engineer",  "ASSUMED billing class",   847.00, 0.327, 0.591, 0.082, 0.424),
    ("Randy Pynenberg",  "Engineer IV",         "org chart",                68.00, 0.897, 0.000, 0.103, 0.950),
    ("Auston Hopson",    "Graphics Technician", "org chart",                18.00, 1.000, 0.000, 0.000, 1.000),
    ("Tom Little",       "Engineer IV",         "org chart — no hours in KPI report", None, None, None, None, None),
    ("Erin Bryden",      "Engineer IV",         "org chart — no hours in KPI report", None, None, None, None, None),
    ("Patrick Nicholson","Senior Designer",     "org chart — no hours in KPI report", None, None, None, None, None),
]
r0 = 5
for i, (name, cls, src, hrs, bil, non, bid, tgt) in enumerate(ROSTER):
    r = r0 + i
    ws.cell(row=r, column=1, value=name).border = BOX
    c2 = ws.cell(row=r, column=2, value=cls)
    c2.border = BOX
    src_cell = ws.cell(row=r, column=3, value=src)
    src_cell.border = BOX
    src_cell.font = F_FLAG if "ASSUMED" in src or "no hours" in src else F_NOTE
    if "ASSUMED" in src:
        c2.font = F_FLAG
    rate_cell = ws.cell(row=r, column=4, value=f"=Rates!B{RATE_ROWS[cls]}")
    rate_cell.number_format = USD
    rate_cell.border = BOX
    for col, val, fmt in [(5, hrs, HRS), (6, bil, PCT), (7, non, PCT), (8, bid, PCT)]:
        c = ws.cell(row=r, column=col, value=val)
        c.number_format = fmt
        c.border = BOX
    bh = ws.cell(row=r, column=9, value=f"=IF(E{r}=\"\",\"\",E{r}*F{r})" if hrs is not None else None)
    bh.number_format = HRS
    bh.border = BOX
    tg = ws.cell(row=r, column=10, value=tgt)
    tg.number_format = PCT
    tg.border = BOX
    var = ws.cell(row=r, column=11, value=(f"=F{r}-J{r}" if hrs is not None else None))
    var.number_format = PCT
    var.border = BOX
    val_c = ws.cell(row=r, column=12, value=(f"=I{r}*D{r}" if hrs is not None else None))
    val_c.number_format = '"$"#,##0'
    val_c.border = BOX
rt = r0 + len(ROSTER)
ws.cell(row=rt, column=1, value="TOTAL / blended").font = F_SUB
for col, fmt in [(5, HRS), (9, HRS), (12, '"$"#,##0')]:
    letter = get_column_letter(col)
    c = ws.cell(row=rt, column=col, value=f"=SUM({letter}{r0}:{letter}{rt-1})")
    c.number_format = fmt
    c.font = F_SUB
b = ws.cell(row=rt, column=4, value=f"=L{rt}/I{rt}")
b.number_format = USD
b.font = F_SUB
for col in range(1, 13):
    ws.cell(row=rt, column=col).fill = FILL_TOT
    ws.cell(row=rt, column=col).border = BOX
ws.cell(row=rt + 2, column=1,
        value=("KPI report totals: 8,052.25 total hrs · 4,747 billable · 2,896 non-billable · 405 bids · 59.0% total "
               "billable utilization. Roster above sums slightly lower (~193 hrs unattributed on the report page).")).font = F_NOTE
ws.cell(row=rt + 3, column=1,
        value="Blended card rate (col D total) is the mix-weighted average bill rate of actual billable hours — used by the Realization tab.").font = F_NOTE
widths(ws, {"A": 18, "B": 20, "C": 30, "D": 11, "E": 12, "F": 11, "G": 13, "H": 9, "I": 12, "J": 12, "K": 12, "L": 16})
ROSTER_TOTAL_ROW = rt

# ============================================================== Baseline P&L
ws = wb.create_sheet("Baseline P&L")
title(ws, "CVR — Baseline P&L, January–May 2026 (all figures $K)",
      "Revenue / Gross Profit / Net Income published in the KPI report. Cost of Sales and Operating Expenses are DERIVED.")
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May"]
header(ws, 4, list(range(1, 8)), ["CVR ACTUALS ($K)"] + MONTHS + ["YTD"])
REV = [110.5, 96.8, 39.3, 94.6, 85.5]
GP = [52.6, 24.2, -15.7, 33.7, -18.0]
NI = [-41.7, -49.1, -98.8, -60.5, -54.5]


def put_row(ws, r, label, values=None, formula=None, fmt=KFMT, font=None, derived=False):
    lab = ws.cell(row=r, column=1, value=label + (" (derived)" if derived else ""))
    if font:
        lab.font = font
    if derived:
        lab.font = F_FLAG
    for i in range(5):
        col = 2 + i
        if values is not None:
            c = ws.cell(row=r, column=col, value=values[i])
        else:
            c = ws.cell(row=r, column=col, value=formula(get_column_letter(col)))
        c.number_format = fmt
        c.border = BOX
    y = ws.cell(row=r, column=7, value=f"=SUM(B{r}:F{r})")
    y.number_format = fmt
    y.border = BOX
    y.fill = FILL_TOT
    return r + 1


r = 5
r = put_row(ws, r, "Revenue", values=REV, font=F_SUB)
r = put_row(ws, r, "Cost of Sales", formula=lambda L: f"={L}5-{L}7", derived=True)
r = put_row(ws, r, "Gross Profit", values=GP, font=F_SUB)
gp_r = r - 1
c = ws.cell(row=r, column=1, value="Gross margin %")
for i in range(5):
    col = get_column_letter(2 + i)
    cc = ws.cell(row=r, column=2 + i, value=f"={col}{gp_r}/{col}5")
    cc.number_format = PCT
cc = ws.cell(row=r, column=7, value=f"=G{gp_r}/G5")
cc.number_format = PCT
r += 1
r = put_row(ws, r, "Operating Expenses", formula=lambda L: f"={L}{gp_r}-{L}{r+1}", derived=True)
ni_row = r
r = put_row(ws, r, "Net Income / (Loss)", values=NI, font=F_SUB)
c = ws.cell(row=r, column=1, value="Net margin %")
for i in range(5):
    col = get_column_letter(2 + i)
    cc = ws.cell(row=r, column=2 + i, value=f"={col}{ni_row}/{col}5")
    cc.number_format = PCT
cc = ws.cell(row=r, column=7, value=f"=G{ni_row}/G5")
cc.number_format = PCT
r += 2

header(ws, r, list(range(1, 8)), ["BUDGET ($K, monthly = Δ of published YTD)"] + MONTHS + ["YTD"])
r += 1
BUD_REV_C = [169.1, 338.2, 507.3, 676.4, 845.5]
BUD_GP_C = [104.1, 208.2, 312.3, 416.5, 520.6]
BUD_NI_C = [11.6, 23.2, 34.9, 46.5, 58.1]


def monthly(cum):
    return [round(cum[0], 1)] + [round(b - a, 1) for a, b in zip(cum, cum[1:])]


bud_rev_row = r
r = put_row(ws, r, "Revenue — budget", values=monthly(BUD_REV_C))
bud_gp_row = r
r = put_row(ws, r, "Gross Profit — budget", values=monthly(BUD_GP_C))
bud_ni_row = r
r = put_row(ws, r, "Net Income — budget", values=monthly(BUD_NI_C))
r += 1
header(ws, r, [1, 2, 3, 4], ["VARIANCE — YTD ($K)", "Actual", "Budget", "Act − Bud"])
r += 1
for label, act_row, bud_row in [("Revenue", 5, bud_rev_row), ("Gross Profit", gp_r, bud_gp_row),
                                ("Net Income", ni_row, bud_ni_row)]:
    ws.cell(row=r, column=1, value=label).border = BOX
    a = ws.cell(row=r, column=2, value=f"=G{act_row}")
    bgt = ws.cell(row=r, column=3, value=f"=G{bud_row}")
    d = ws.cell(row=r, column=4, value=f"=B{r}-C{r}")
    for cell in (a, bgt, d):
        cell.number_format = KFMT
        cell.border = BOX
    r += 1
ws.cell(row=r + 1, column=1,
        value=("Reading: YTD revenue at 50% of budget; two negative-GP months (Mar, May); OpEx derived at ~$76K/mo "
               "average. Phase 2 replaces derived lines with USC's actual line items.")).font = F_NOTE
widths(ws, {"A": 34, "B": 10, "C": 10, "D": 10, "E": 10, "F": 10, "G": 11})

# =========================================================== HWC EPC Baseline
ws = wb.create_sheet("HWC EPC Baseline")
title(ws, "HWC EPC — Baseline, January–May 2026 (all figures $K)",
      "Context for the engineering P&L: design quality and constructability directly affect these outcomes.")
header(ws, 4, list(range(1, 8)), ["HWC EPC ($K)"] + MONTHS + ["YTD"])
r = 5
r = put_row(ws, r, "Revenue", values=[732.2, 207.7, -321.0, -125.1, 678.6], font=F_SUB)
r = put_row(ws, r, "Net Profit / (Loss)", values=[-318.3, -346.9, -672.8, -656.5, -114.8], font=F_SUB)
c = ws.cell(row=r, column=1, value="Net margin % (YTD)")
cc = ws.cell(row=r, column=7, value="=G6/G5")
cc.number_format = PCT
r += 2
for note in [
    "Negative revenue in March and April are billing adjustments/write-downs per the KPI report charts.",
    "Active EPC scorecard jobs (7/11/2026): 4 substation projects, PO $4.3M–$6.6M, plan margins 12.7%–16.4%,",
    "cost variance +3.6% to −19.8%. Engineering share of paired EPC pursuits runs ~4–5% of EPC job value.",
]:
    ws.cell(row=r, column=1, value=note).font = F_NOTE
    r += 1
widths(ws, {"A": 30, "B": 10, "C": 10, "D": 10, "E": 10, "F": 10, "G": 11})

# ================================================================ Realization
ws = wb.create_sheet("Realization")
title(ws, "Realization Analysis — where does the rate card go?",
      "Compares what YTD billable hours are worth at card rates against revenue actually recognized.")
header(ws, 4, [1, 2, 3], ["Step", "Value", "Source / formula"])
real_rows = [
    ("Billable hours YTD (through 7/11)", "='Roster & Hours'!I%d" % ROSTER_TOTAL_ROW, HRS,
     "Roster & Hours total (KPI report: 4,747)"),
    ("Adjustment factor to Jan–May", "=Assumptions!B8", '0.000', "Assumptions — scales hours to match Jan–May financials"),
    ("Est. billable hours, Jan–May", "=B5*B6", HRS, "calc"),
    ("Blended card rate ($/hr)", "='Roster & Hours'!D%d" % ROSTER_TOTAL_ROW, USD, "mix-weighted from actual billable hours"),
    ("Value at card rates, Jan–May ($K)", "=B7*B8/1000", KFMT, "calc"),
    ("Actual revenue, Jan–May ($K)", "='Baseline P&L'!G5", KFMT, "Baseline P&L"),
    ("Realization %", "=B10/B9", PCT, "actual ÷ card value"),
    ("Implied realized rate ($/hr)", "=B10*1000/B7", USD, "calc"),
    ("Unexplained gap ($K)", "=B9-B10", KFMT, "to be decomposed below"),
]
r = 5
for label, formula, fmt, src in real_rows:
    ws.cell(row=r, column=1, value=label).border = BOX
    c = ws.cell(row=r, column=2, value=formula)
    c.number_format = fmt
    c.border = BOX
    s = ws.cell(row=r, column=3, value=src)
    s.font = F_NOTE
    s.border = BOX
    r += 1
r += 1
ws.cell(row=r, column=1, value="Gap decomposition — fill after the finance review (must sum to the unexplained gap):").font = F_SUB
r += 1
for bucket in ["Unbilled WIP / billing lag ($K)", "Write-offs & discounts ($K)",
               "Internal work at reduced transfer rates ($K)", "Other / unreconciled ($K)"]:
    ws.cell(row=r, column=1, value=bucket).border = BOX
    c = ws.cell(row=r, column=2, value=0)
    c.font = F_INPUT
    c.number_format = KFMT
    c.border = BOX
    r += 1
chk = ws.cell(row=r, column=1, value="Decomposition vs gap")
c = ws.cell(row=r, column=2, value=f"=SUM(B{r-4}:B{r-1})-B13")
c.number_format = KFMT
ws.cell(row=r, column=3, value="should be 0 when fully explained").font = F_NOTE
r += 3

ws.cell(row=r, column=1, value="Annualized revenue grid ($K) — current roster, utilization × realization").font = F_SUB
r += 1
grid_top = r
utils = [0.59, 0.65, 0.70, 0.75]
reals = [0.60, 0.75, 0.90, 1.00]
ws.cell(row=grid_top, column=1, value="Utilization ↓ / Realization →").font = F_NOTE
for j, real in enumerate(reals):
    c = ws.cell(row=grid_top, column=2 + j, value=real)
    c.number_format = PCT
    c.font = F_H
    c.fill = FILL_H
    c.border = BOX
for i, u in enumerate(utils):
    rr = grid_top + 1 + i
    c = ws.cell(row=rr, column=1, value=u)
    c.number_format = PCT
    c.font = F_H
    c.fill = FILL_H
    c.border = BOX
    for j in range(len(reals)):
        col = get_column_letter(2 + j)
        cell = ws.cell(
            row=rr, column=2 + j,
            value=(f"=Assumptions!B7*Assumptions!B5*$A{rr}*{col}${grid_top}*$B$8/1000"))
        cell.number_format = '#,##0'
        cell.border = BOX
r = grid_top + len(utils) + 2
for note in [
    "Grid = active billable FTEs × hrs/FTE-yr × utilization × realization × blended card rate, in $K.",
    "Bottom-right cell (75% util, 100% realization) shows current-roster ceiling; compare against the $6,000K Year-2 goal",
    "to size the hiring plan (Phase 2). Timing caveat: hours run through 7/11 while revenue runs through 5/31 —",
    "the adjustment factor on Assumptions handles this; refine when finance provides hours-matched revenue.",
]:
    ws.cell(row=r, column=1, value=note).font = F_NOTE
    r += 1
widths(ws, {"A": 44, "B": 13, "C": 13, "D": 13, "E": 13})

# ======================================================== Checks & Open Items
ws = wb.create_sheet("Checks & Open Items")
title(ws, "Tie-outs to the KPI report and open questions")
header(ws, 4, [1, 2, 3, 4], ["Check", "Model", "KPI report", "Status"])
checks = [
    ("CVR revenue YTD ($K)", "='Baseline P&L'!G5", 426.7),
    ("CVR gross profit YTD ($K)", "='Baseline P&L'!G7", 76.8),
    ("CVR net income YTD ($K)", "='Baseline P&L'!G10", -304.7),
    ("EPC revenue YTD ($K)", "='HWC EPC Baseline'!G5", 1172.4),
    ("EPC net profit YTD ($K)", "='HWC EPC Baseline'!G6", -2109.3),
    ("Billable hours YTD", "='Roster & Hours'!I%d" % ROSTER_TOTAL_ROW, 4747.0),
]
r = 5
for label, formula, target in checks:
    ws.cell(row=r, column=1, value=label).border = BOX
    m = ws.cell(row=r, column=2, value=formula)
    m.number_format = KFMT
    m.border = BOX
    t = ws.cell(row=r, column=3, value=target)
    t.number_format = KFMT
    t.border = BOX
    tol = 60 if "hours" in label.lower() else 1.0
    s = ws.cell(row=r, column=4, value=f'=IF(ABS(B{r}-C{r})<={tol},"TIES","CHECK")')
    s.border = BOX
    r += 1
ws.cell(row=r, column=1,
        value="Billable-hours tolerance is wider: ~193 hrs on the KPI page are not attributed to a named employee.").font = F_NOTE
r += 2
ws.cell(row=r, column=1, value="OPEN ITEMS (Phase 1 exit criteria)").font = F_SUB
r += 1
items = [
    "Finance: fringe/benefit load %, salaries for loaded-cost-vs-bill-rate margin by role.",
    "Finance: USC corporate overhead allocation methodology and current CVR allocation.",
    "Finance: interest expense basis for OILI.",
    "Finance: transfer-pricing treatment of CVR hours on USC/HWC jobs — primary suspect for the realization gap.",
    "Confirm billing classifications for Taylor Nelson and Kevin Metts (not on the org chart).",
    "Confirm billing classification used for VP Engineering time (assumed Principal Engineer).",
    "Tom Little, Erin Bryden, Patrick Nicholson show no hours on the KPI utilization page — confirm where their time goes.",
    "Decompose the realization gap (Realization tab) once finance data arrives.",
    "Confirm the 2026 CVR budget behind the KPI report's budget lines.",
]
for item in items:
    ws.cell(row=r, column=1, value="• " + item)
    r += 1
widths(ws, {"A": 90, "B": 12, "C": 12, "D": 10})

import os
os.makedirs("model", exist_ok=True)
out = "model/CVR_Model_Phase1_Baseline.xlsx"
wb.save(out)
print(f"wrote {out}")
